import csv, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
# Names are real Indian jewellers; scale/source columns carry only what research found. Contact columns stay blank on purpose (enrich from public business pages; do not guess).
H=["ID","Business","HQ city","State","Segment","Reported scale (approx, verify)","Listed","Website (verified only)","Budget fit (<= Rs 2,00,000 per model)","Suggested plan","Outreach angle","Researched this session","Source","Target contact role","Contact name","Email","Phone","Status","Next action"]
S1="Indian Jeweller / Morningstar"
rows=[
("Manoj Vaibhav Gems N Jewellers","Visakhapatnam","Andhra Pradesh","Regional chain","13 showrooms in AP & Telangana (incl. 2 franchise)","NSE/BSE","","Core","Growth","Franchise + multi-showroom enquiries in one CRM","Yes","https://dfd.dr.morningstar.com.au/investments/security/NSE/MVGJL"),
("Thangamayil Jewellery","Madurai","Tamil Nadu","Regional chain","~54-66 stores (sources differ)","NSE/BSE","thangamayil.com","Growth/Pro","Pro","Tier-2/3 expansion: consistent product visuals and lead follow-up across stores","Yes","https://indianretailer.com/franchisetv/archives/news/thangamayil-jewellery-opens-two-stores-chennai-expands-retail-network-66-stores"),
("PNG Jewellers (P N Gadgil & Sons)","Pune","Maharashtra","Regional chain","~65 stores reported Dec 2025","","","Growth/Pro","Pro","Campaign visuals + approval workflow for many stores","Yes","https://www.indianjeweller.in/Indian-Jewellery-News/15120/index.html"),
("P N Gadgil Jewellers Ltd (PNGS)","Pune","Maharashtra","Regional chain","34-68 stores (sources differ); Maharashtra, Karnataka, Gujarat","NSE (PNGJL)","","Growth/Pro","Pro","Separate company from PNG Jewellers above; confirm which entity before outreach","Yes","https://www.investywise.com/?p=58933"),
("Kataria Jewellers","Indore","Madhya Pradesh","Regional","Flagship 6,000 sq ft showroom opened Nov 2025","","","Core","Growth","New flagship: capture walk-in and online enquiries","Yes","https://mail.indianjeweller.in/search-news/Madhya Pradesh"),
("Manepally Jewellers","Hyderabad","Telangana","Family jeweller","Family business, ~125 years in Hyderabad (press)","","","Core","Starter/Growth","Heritage designs: catalogue + 360 view for bridal enquiries","Yes","https://deccanchronicle.com/lifestyle/viral-and-trending/060218/a-glittering-family-legacy.html"),
("Gem Palace","Jaipur","Rajasthan","Luxury family jeweller","Family-run, court jewellers to Jaipur royals (press)","","","Core","Growth","High-touch clients: private-viewing enquiries and follow-ups","Yes","https://www.scmp.com/magazines/style/luxury/jewellery/article/3370805/story-jaipur-jeweller-gem-palace-it-debuts-lane-crawford-hong-kong"),
("Surana Jewellers (Bhuramal Rajmal Surana)","Jaipur","Rajasthan","Family jeweller","Self-described since 1735 (own blog; verify)","","","Core","Starter/Growth","Jadau/Kundan catalogue and enquiry tracking","Yes","https://suranajewellery.com/tag/best-jewellers-in-jaipur/"),
("Khazanchi Jewellers","Chennai","Tamil Nadu","Regional","Incorporated 1996","BSE","","Core/Growth","Growth","Listed SME: visuals and CRM without a big IT team","Yes","https://equitymaster.com/share-price/KHAN/KHAZANCHI-JEWELLERS-Share-Price"),
("Motisons Jewellers","Jaipur","Rajasthan","Regional","Operating since 1998","Listed (verify exchange)","","Core/Growth","Growth","Jaipur showroom enquiries and festive campaigns","Yes","https://indiainfoline.com/company/motisons-jewellers-ltd/summary"),
("Shringar House of Mangalsutra","Mumbai","Maharashtra","B2B manufacturer","Mangalsutra maker for B2B clients","NSE (SHRINGARMS)","","Core (B2B angle)","Growth","B2B catalogue sharing and dealer-enquiry CRM","Yes","https://www.morningstar.com.au/investments/security/NSE/SHRINGARMS"),
("RBZ Jewellers","Ahmedabad","Gujarat","Regional","Named as listed peer in IPO comparison","Listed (verify)","","Core/Growth","Growth","Wedding-season lead capture and follow-up","Partly","https://ipocentral.in/shringar-house-vs-rbz-vs-sky-gold-vs-utssav-cz/?rand=12352"),
("TBZ - The Original (Tribhovandas Bhimji Zaveri)","Mumbai","Maharashtra","Heritage chain","Heritage brand founded 1864 (press)","Listed (verify)","","Growth/Pro","Pro","Heritage brand: consistent visuals and approvals","Partly","https://www.indianretailer.com/article/retail-business/retail/top-5-indian-jewellery-brands-strongest-retail-footprint"),
("Senco Gold & Diamonds","Kolkata","West Bengal","National chain","150-164 showrooms (sources differ)","NSE (SENCO)","sencogoldanddiamonds.com","Enterprise (above Rs 2L; later)","Pro (pilot) / custom","Pilot for one city or one collection, then upgrade","Yes","https://www.indianretailer.com/news/senco-gold-and-diamonds-celebrates-milestone-150-showrooms-india-vidya-balan"),
("PC Jeweller","Delhi","Delhi NCR","National chain","North India, mall-based","NSE/BSE","","Enterprise (above Rs 2L; later)","custom","Campaign visuals and lead follow-up pilot","Yes","https://corporate.pcjeweller.com/?p=13"),
# Known names, not researched this session: verify before outreach
("Waman Hari Pethe Sons","Pune","Maharashtra","Heritage regional","Not researched","","","Growth/Pro","Growth/Pro","Heritage Pune jeweller: bridal enquiries and campaigns","No - verify",""),
("Popley & Sons","Delhi","Delhi NCR","Family jeweller","Not researched","","","Core","Growth","Delhi family jeweller: enquiry tracking","No - verify",""),
("Amrapali Jewels","Jaipur","Rajasthan","Designer jeweller","Not researched","","","Core","Growth","Designer catalogue + 360 view","No - verify",""),
("Kishandas & Co","Hyderabad","Telangana","Family jeweller","Not researched","","","Core","Starter/Growth","Catalogue and enquiry CRM","No - verify",""),
("Lalithaa Jewellery","Chennai","Tamil Nadu","Regional chain","Not researched","","","Growth/Pro","Growth/Pro","Multi-store follow-up and visuals","No - verify",""),
("Vummidi Bangaru Jewellers","Chennai","Tamil Nadu","Regional","Not researched","","","Core/Growth","Growth","Chennai showroom enquiries","No - verify",""),
("GRT Jewellers","Chennai","Tamil Nadu","Regional chain","Not researched","","","Enterprise (verify size)","custom","Pilot","No - verify",""),
("Bhima Jewellery","Thiruvananthapuram","Kerala","Regional chain","Not researched","","","Enterprise (verify size)","custom","Pilot","No - verify",""),
("Jos Alukkas","Thrissur","Kerala","Regional chain","Not researched","","","Growth/Pro","Pro","Kerala bridal campaigns and leads","No - verify",""),
("Chemmanur International Jewellers","Kochi","Kerala","Regional chain","Not researched","","","Growth/Pro","Pro","Kerala multi-store lead follow-up","No - verify",""),
]
out=[]
for i,r in enumerate(rows,1):
    biz,city,st,seg,scale,listed,web,fit,plan,angle,res,src=r
    role="Owner / Managing Director; Head of Marketing; E-commerce or IT head"
    out.append([f"L{i:03d}",biz,city,st,seg,scale,listed,web,fit,plan,angle,res,src,role,"","","","New","Find official business contact from the brand's own website/Contact page; send 1 intro email + demo link"])
with open('india-jeweller-leads.csv','w',newline='',encoding='utf-8') as f:
    w=csv.writer(f);w.writerow(H);w.writerows(out)
wb=openpyxl.Workbook();ws=wb.active;ws.title="Leads"
ws.append(H)
for r in out:ws.append(r)
hf=PatternFill('solid',fgColor='461318')
for c in ws[1]:
    c.font=Font(bold=True,color='F8EEDB');c.fill=hf;c.alignment=Alignment(wrap_text=True,vertical='center')
widths=[7,34,16,16,18,40,14,24,24,16,48,14,50,34,18,24,16,10,48]
for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w
for row in ws.iter_rows(min_row=2):
    for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
p=wb.create_sheet("Plans (proposed)")
p.append(["Plan (per model / licence)","Proposed one-time price (INR)","Best for","Included","Upgrade path"])
plans=[
("Starter",49999,"1 showroom, <=1,000 designs","Catalogue + price PDF import, public enquiry page, AI lead scoring, Kanban CRM, role logins","Pay the difference to move up"),
("Growth",99999,"1-10 showrooms","Starter + Creative Studio (approval-gated), 360 view, branded e-mail approvals, analytics, showcase PDF","Pay the difference to Pro"),
("Pro",199999,"10+ showrooms or multi-brand","Growth + multi-store roles, custom branding and domain, onboarding and training, priority support","Custom enterprise scope above Rs 2,00,000"),
]
for x in plans:p.append(list(x))
p.append([]);p.append(["Note: prices are a proposal to set your own numbers. Your ceiling was read as Rs 2,00,000 (Rs 2 lakh) per model. Annual support/hosting is not included; decide separately."])
for c in p[1]:c.font=Font(bold=True,color='F8EEDB');c.fill=hf
for i,w in enumerate([26,26,34,80,36],1):p.column_dimensions[get_column_letter(i)].width=w
n=wb.create_sheet("Read me")
for line in ["How this list was built: 15 rows researched on the public web this session (Source column); 10 more are well-known jewellers added by name only and marked 'No - verify'.","Contact name, e-mail and phone are blank on purpose. Take them only from the brand's own website or public business listing; never guess or buy scraped personal data.","Store counts differ between sources and change often. Treat them as approximate.","Outreach rules (India): business e-mails need a clear sender, an opt-out line and honest subject lines; do not send bulk SMS/WhatsApp without consent (DPDP Act 2023, TRAI DLT rules).","Budget fit: Enterprise rows exceed Rs 2,00,000; approach with a small pilot only.","This is a starter list of 25, not every jeweller in India (there are hundreds of thousands). To grow it: GJEPC and IBJA member directories, Google Maps by city, Indian Jeweller news, trade shows (IIJS)."]:
    n.append([line])
n.column_dimensions['A'].width=140
wb.save('india-jeweller-leads.xlsx');print(len(out),'rows')
